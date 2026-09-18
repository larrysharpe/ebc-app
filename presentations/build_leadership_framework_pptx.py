#!/usr/bin/env python3
"""
Ebenezer Baptist Church — Condensed leadership briefing.

Communications & Digital Engagement Workstream
Building upon the Leadership Conference objectives
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OFFICIAL_LOGO = ASSETS / "ebc-logo.png"
DARK_LOGO = ASSETS / "ebc-logo-dark.png"
OUTPUT = ROOT / "EBC_Communications_Workstream_Briefing.pptx"

PRES_TITLE = "Communications & Digital Engagement Workstream"
PRES_SUBTITLE = "Building Upon the Leadership Conference"

# Colors sampled from official logo
BURGUNDY = RGBColor(0x80, 0x20, 0x20)       # #802020
BURGUNDY_DARK = RGBColor(0x50, 0x14, 0x18)  # deep companion
GOLD = RGBColor(0xA0, 0x80, 0x20)           # #A08020
GOLD_SOFT = RGBColor(0xC4, 0xA8, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF2, 0xF0)
MID_GRAY = RGBColor(0xD8, 0xD2, 0xCC)
CHARCOAL = RGBColor(0x2A, 0x2A, 0x2A)
SOFT_CHARCOAL = RGBColor(0x5C, 0x55, 0x52)
CARD_BORDER = RGBColor(0xE2, 0xDC, 0xD6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_SLIDES = 10


def _set_run_font(run, *, size: Pt, bold: bool = False, color: RGBColor = CHARCOAL, name: str = "Calibri") -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _set_tf_margins(tf, margin: float = 0.08) -> None:
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.6)
    tf.margin_bottom = Inches(margin * 0.6)


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = CHARCOAL,
    align=PP_ALIGN.LEFT,
    font_name: str = "Calibri",
    anchor=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=Pt(size), bold=bold, color=color, name=font_name)


def add_rect(slide, left, top, width, height, fill: RGBColor, *, line: RGBColor | None = None, line_pt: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def add_round_rect(slide, left, top, width, height, fill: RGBColor, *, line: RGBColor | None = None, line_pt: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def shape_text(shape, text: str, *, size: int = 12, bold: bool = False, color: RGBColor = WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    _set_tf_margins(tf)
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=Pt(size), bold=bold, color=color)


def shape_multiline(shape, lines: Sequence[Tuple[str, dict]], *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    _set_tf_margins(tf, 0.12)
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    for i, (text, opts) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opts.get("space_after", 2))
        run = p.add_run()
        run.text = text
        _set_run_font(
            run,
            size=Pt(opts.get("size", 12)),
            bold=opts.get("bold", False),
            color=opts.get("color", WHITE),
        )


def add_logo(slide, *, left=Inches(0.4), top=Inches(0.22), height=Inches(0.58), dark: bool = False) -> None:
    path = DARK_LOGO if dark and DARK_LOGO.exists() else OFFICIAL_LOGO
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, height=height)


def add_footer(slide, page: int) -> None:
    add_rect(slide, Inches(0.45), Inches(7.08), Inches(12.43), Pt(1.5), GOLD)
    add_textbox(slide, Inches(0.45), Inches(7.15), Inches(9.2), Inches(0.25), PRES_TITLE, size=9, color=SOFT_CHARCOAL)
    add_textbox(
        slide,
        Inches(9.8),
        Inches(7.15),
        Inches(3.1),
        Inches(0.25),
        f"Ebenezer Baptist Church  |  {page}",
        size=9,
        color=SOFT_CHARCOAL,
        align=PP_ALIGN.RIGHT,
    )


def add_slide_header(slide, title: str, eyebrow: str | None = None) -> None:
    add_logo(slide, left=Inches(0.4), top=Inches(0.2), height=Inches(0.55))
    title_left = Inches(3.35)
    if eyebrow:
        add_textbox(slide, title_left, Inches(0.2), Inches(9.4), Inches(0.25), eyebrow.upper(), size=10, bold=True, color=GOLD)
        add_textbox(slide, title_left, Inches(0.42), Inches(9.4), Inches(0.45), title, size=26, bold=True, color=BURGUNDY)
    else:
        add_textbox(slide, title_left, Inches(0.28), Inches(9.4), Inches(0.5), title, size=26, bold=True, color=BURGUNDY)
    add_rect(slide, title_left, Inches(0.95), Inches(1.3), Pt(3), BURGUNDY)


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------------------
# Slides (10 total — condensed for leadership consumption)
# ---------------------------------------------------------------------------

def slide_01_title(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, BURGUNDY_DARK)
    add_rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, GOLD)

    add_logo(slide, left=Inches(0.7), top=Inches(1.4), height=Inches(1.15))

    add_textbox(slide, Inches(0.7), Inches(2.9), Inches(11.5), Inches(0.35), "LEADERSHIP BRIEFING", size=12, bold=True, color=GOLD)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(3.35),
        Inches(11.8),
        Inches(1.2),
        "Communications &\nDigital Engagement Workstream",
        size=34,
        bold=True,
        color=WHITE,
    )
    add_rect(slide, Inches(0.7), Inches(4.75), Inches(1.5), Pt(2.5), GOLD)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(5.0),
        Inches(11),
        Inches(0.6),
        "Building upon the Leadership Conference\nto strengthen execution—not reorganize ministries",
        size=16,
        color=GOLD_SOFT,
    )
    add_textbox(slide, Inches(0.7), Inches(6.6), Inches(11), Inches(0.35), "Ebenezer Baptist Church", size=12, color=MID_GRAY)

    set_notes(
        slide,
        """Open by affirming the Leadership Conference. This briefing is about coordinating communications and digital work already underway—not creating a new ministry or restructuring.
Keep the room focused: one workstream, clear ask, short discussion.""",
    )


def slide_02_the_point(prs: Presentation) -> None:
    """Lead with the bottom line — leaders skim."""
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "The Point", "One-Page Summary")
    add_footer(slide, 2)

    left = add_round_rect(slide, Inches(0.5), Inches(1.25), Inches(6.0), Inches(5.3), LIGHT_GRAY, line=CARD_BORDER)
    shape_multiline(
        left,
        [
            ("What we are proposing", {"size": 12, "bold": True, "color": GOLD, "space_after": 10}),
            (
                "Create an Operational Workstream that coordinates three related Leadership Conference objectives so they move together with shared planning, clear ownership, and measurable results.",
                {"size": 15, "color": CHARCOAL, "space_after": 16},
            ),
            ("What this is not", {"size": 12, "bold": True, "color": GOLD, "space_after": 8}),
            ("Not a reorganization", {"size": 14, "bold": True, "color": BURGUNDY, "space_after": 4}),
            ("Not a new ministry", {"size": 14, "bold": True, "color": BURGUNDY, "space_after": 4}),
            ("Not a bylaw change", {"size": 14, "bold": True, "color": BURGUNDY, "space_after": 0}),
        ],
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,
    )

    asks = [
        ("1", "Endorse", "Adopt this workstream as the way we execute these Conference objectives."),
        ("2", "Name a Lead", "Identify an Operational Lead in the next 90 days."),
        ("3", "Review Quarterly", "Receive a simple progress update each quarter."),
    ]
    for i, (num, title, desc) in enumerate(asks):
        y = Inches(1.25) + i * Inches(1.7)
        card = add_round_rect(slide, Inches(6.8), y, Inches(5.9), Inches(1.5), WHITE, line=CARD_BORDER)
        badge = add_round_rect(slide, Inches(7.05), y + Inches(0.4), Inches(0.7), Inches(0.7), BURGUNDY)
        shape_text(badge, num, size=18, bold=True, color=WHITE)
        add_textbox(slide, Inches(8.0), y + Inches(0.3), Inches(4.4), Inches(0.4), title, size=18, bold=True, color=BURGUNDY)
        add_textbox(slide, Inches(8.0), y + Inches(0.75), Inches(4.4), Inches(0.55), desc, size=13, color=SOFT_CHARCOAL)

    set_notes(
        slide,
        """Lead with this slide. Many leaders only need this page.
If time is short, jump from here to Recommendations and Decision.
Emphasize the three “nots” early to reduce anxiety.""",
    )


def slide_03_build_upon(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Building Upon the Conference", "Vision → Execution")
    add_footer(slide, 3)

    stages = [
        ("Leadership\nConference", "Vision set"),
        ("Related\nObjectives", "Priorities named"),
        ("This\nWorkstream", "Execution coordinated"),
    ]
    for i, (title, sub) in enumerate(stages):
        x = Inches(0.7) + i * Inches(4.15)
        fill = GOLD if i == 2 else BURGUNDY
        text_c = BURGUNDY_DARK if i == 2 else WHITE
        box = add_round_rect(slide, x, Inches(1.8), Inches(3.7), Inches(2.0), fill)
        shape_multiline(
            box,
            [
                (title, {"size": 18, "bold": True, "color": text_c, "space_after": 8}),
                (sub, {"size": 13, "color": text_c, "space_after": 0}),
            ],
        )
        if i < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(3.8), Inches(2.45), Inches(0.3), Inches(0.45))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GOLD
            arrow.line.fill.background()

    banner = add_round_rect(slide, Inches(0.7), Inches(4.3), Inches(11.9), Inches(2.15), LIGHT_GRAY, line=CARD_BORDER)
    shape_multiline(
        banner,
        [
            ("Focus: better coordination and execution", {"size": 18, "bold": True, "color": BURGUNDY, "space_after": 10}),
            (
                "Every Leadership Conference objective remains intact. This workstream helps related communications and digital initiatives plan together, communicate clearly, and show measurable progress.",
                {"size": 14, "color": CHARCOAL, "space_after": 0},
            ),
        ],
        align=PP_ALIGN.CENTER,
    )

    set_notes(
        slide,
        """One diagram only. Conference work was excellent; we are operationalizing it.
Gold box = coordination layer, not a new org box on a chart.""",
    )


def slide_04_why(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Why a Workstream?", "Coordinate What Already Exists")
    add_footer(slide, 4)

    benefits = [
        ("Shared Planning", "One calendar, aligned priorities"),
        ("Clear Ownership", "Someone accountable for cadence"),
        ("Consistent Voice", "Messaging that supports every ministry"),
        ("Less Duplication", "Fewer parallel efforts"),
        ("Measurable Progress", "Simple metrics leadership can review"),
        ("Sustainable Pace", "Quarterly rhythm, not one-time push"),
    ]
    for i, (t, d) in enumerate(benefits):
        r, c = i // 3, i % 3
        x = Inches(0.5) + c * Inches(4.2)
        y = Inches(1.35) + r * Inches(2.5)
        card = add_round_rect(slide, x, y, Inches(4.0), Inches(2.2), WHITE, line=CARD_BORDER)
        add_rect(slide, x, y, Inches(4.0), Inches(0.12), BURGUNDY if i % 2 == 0 else GOLD)
        add_textbox(slide, x + Inches(0.25), y + Inches(0.5), Inches(3.5), Inches(0.5), t, size=18, bold=True, color=BURGUNDY, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.3), y + Inches(1.15), Inches(3.4), Inches(0.7), d, size=14, color=SOFT_CHARCOAL, align=PP_ALIGN.CENTER)

    set_notes(
        slide,
        """Six benefits, plain language. No jargon beyond “workstream.”
Define workstream once: a coordination group for related initiatives—not a ministry.""",
    )


def slide_05_workstream(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Communications Workstream", "Mission & Conference Objectives")
    add_footer(slide, 5)

    mission = add_round_rect(slide, Inches(0.5), Inches(1.25), Inches(12.3), Inches(1.5), BURGUNDY)
    shape_multiline(
        mission,
        [
            ("MISSION", {"size": 11, "bold": True, "color": GOLD, "space_after": 6}),
            (
                "Coordinate communication and digital engagement initiatives to strengthen visibility, consistency, outreach, and support for every ministry.",
                {"size": 16, "color": WHITE, "space_after": 0},
            ),
        ],
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    add_textbox(slide, Inches(0.5), Inches(3.0), Inches(12), Inches(0.35), "Supports these Leadership Conference objectives", size=13, bold=True, color=CHARCOAL)

    objs = [
        ("Digital Ministry &\nInnovation Team", "Innovation capacity"),
        ("E-Ministry &\nDigital Campus", "Digital presence"),
        ("Media &\nCommunity Outreach", "Public witness"),
    ]
    for i, (title, sub) in enumerate(objs):
        x = Inches(0.5) + i * Inches(4.2)
        card = add_round_rect(slide, x, Inches(3.45), Inches(4.0), Inches(2.0), WHITE, line=CARD_BORDER)
        add_rect(slide, x, Inches(3.45), Inches(4.0), Inches(0.1), GOLD)
        shape_multiline(
            card,
            [
                (title, {"size": 16, "bold": True, "color": BURGUNDY, "space_after": 8}),
                (sub, {"size": 13, "color": SOFT_CHARCOAL, "space_after": 0}),
            ],
        )

    set_notes(
        slide,
        """Read the mission. Point to the three Conference objectives—coordinated, not replaced.
This is the only workstream in this briefing; others can be separate presentations later.""",
    )


def slide_06_shared_service(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Supports Every Ministry", "Shared Service Model")
    add_footer(slide, 6)

    hub = add_round_rect(slide, Inches(4.55), Inches(3.15), Inches(4.2), Inches(1.5), BURGUNDY)
    shape_multiline(
        hub,
        [
            ("Communications &\nDigital Engagement", {"size": 15, "bold": True, "color": WHITE, "space_after": 4}),
            ("Workstream", {"size": 12, "color": GOLD, "space_after": 0}),
        ],
    )

    ministries = [
        ("Pastor", Inches(0.55), Inches(1.35)),
        ("Youth", Inches(3.0), Inches(1.35)),
        ("Women", Inches(5.45), Inches(1.35)),
        ("Men", Inches(7.9), Inches(1.35)),
        ("Christian Ed.", Inches(10.35), Inches(1.35)),
        ("Music", Inches(0.55), Inches(5.4)),
        ("Missionary", Inches(3.0), Inches(5.4)),
        ("Outreach", Inches(5.45), Inches(5.4)),
        ("Admin", Inches(7.9), Inches(5.4)),
        ("Media", Inches(10.35), Inches(5.4)),
    ]
    for name, x, y in ministries:
        node = add_round_rect(slide, x, y, Inches(2.25), Inches(0.7), WHITE, line=BURGUNDY, line_pt=1.25)
        shape_text(node, name, size=13, bold=True, color=BURGUNDY)

    banner = add_round_rect(slide, Inches(2.0), Inches(6.3), Inches(9.3), Inches(0.5), GOLD)
    shape_text(banner, "Supporting ministries — not replacing any ministry", size=13, bold=True, color=BURGUNDY_DARK)

    set_notes(
        slide,
        """Key visual for ministry leaders. The workstream amplifies their voice; it does not oversee their ministry.
Repeat the banner sentence.""",
    )


def slide_07_outcomes_metrics(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Outcomes & Metrics", "What Success Looks Like")
    add_footer(slide, 7)

    # Four outcome pillars — compact
    outcomes = [
        ("Communication", "Consistent messaging\nShared calendar"),
        ("Digital", "Website · Social\nLivestream quality"),
        ("Creative", "Photo · Video\nGraphics · Signage"),
        ("Innovation", "AI tools · Forms\nDashboards"),
    ]
    for i, (t, d) in enumerate(outcomes):
        x = Inches(0.5) + i * Inches(3.2)
        card = add_round_rect(slide, x, Inches(1.25), Inches(3.05), Inches(2.0), WHITE, line=CARD_BORDER)
        add_rect(slide, x, Inches(1.25), Inches(3.05), Inches(0.45), BURGUNDY)
        add_textbox(slide, x + Inches(0.1), Inches(1.32), Inches(2.85), Inches(0.35), t, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.15), Inches(1.9), Inches(2.75), Inches(1.1), d, size=13, color=CHARCOAL, align=PP_ALIGN.CENTER)

    add_textbox(slide, Inches(0.5), Inches(3.5), Inches(12), Inches(0.35), "Leadership review metrics (baseline → target)", size=13, bold=True, color=CHARCOAL)

    metrics = [
        "Website Visitors",
        "Social Reach",
        "Livestream Attendance",
        "First-Time Guests",
        "Volunteer Growth",
        "Requests Completed",
    ]
    for i, m in enumerate(metrics):
        r, c = i // 3, i % 3
        x = Inches(0.5) + c * Inches(4.2)
        y = Inches(3.95) + r * Inches(1.25)
        card = add_round_rect(slide, x, y, Inches(4.0), Inches(1.1), LIGHT_GRAY, line=CARD_BORDER)
        add_rect(slide, x, y, Pt(5), Inches(1.1), GOLD)
        shape_text(card, m, size=15, bold=True, color=BURGUNDY)

    set_notes(
        slide,
        """Keep metrics few and understandable. Targets set with the Operational Lead—not tonight.
Avoid vanity-only metrics; guests and completed requests matter to ministry.""",
    )


def slide_08_how_we_run(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "How We Run It", "Ownership & 12-Month Cadence")
    add_footer(slide, 8)

    # Ownership — 3 roles only
    roles = [
        ("Executive Sponsor", "Champions the workstream;\nremoves barriers"),
        ("Operational Lead", "Runs planning cadence;\ncoordinates partners"),
        ("Supporting Ministries", "Bring needs & content;\nshare in the work"),
    ]
    for i, (t, d) in enumerate(roles):
        x = Inches(0.5) + i * Inches(4.2)
        card = add_round_rect(slide, x, Inches(1.25), Inches(4.0), Inches(2.0), WHITE, line=CARD_BORDER)
        add_rect(slide, x, Inches(1.25), Inches(4.0), Inches(0.1), GOLD)
        add_textbox(slide, x + Inches(0.2), Inches(1.55), Inches(3.6), Inches(0.45), t, size=16, bold=True, color=BURGUNDY, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.25), Inches(2.15), Inches(3.5), Inches(0.85), d, size=13, color=SOFT_CHARCOAL, align=PP_ALIGN.CENTER)

    # Timeline
    add_rect(slide, Inches(0.7), Inches(4.55), Inches(11.9), Pt(3), GOLD)
    quarters = [
        ("Q1", "Name Lead"),
        ("Q2", "Annual Plan"),
        ("Q3", "KPI Dashboard"),
        ("Q4", "Review & Improve"),
    ]
    for i, (q, action) in enumerate(quarters):
        x = Inches(1.0) + i * Inches(3.05)
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.85), Inches(4.35), Inches(0.4), Inches(0.4))
        marker.fill.solid()
        marker.fill.fore_color.rgb = BURGUNDY
        marker.line.fill.background()
        shape_text(marker, str(i + 1), size=11, bold=True, color=WHITE)
        card = add_round_rect(slide, x, Inches(5.0), Inches(2.6), Inches(1.35), WHITE, line=CARD_BORDER)
        shape_multiline(
            card,
            [
                (q, {"size": 14, "bold": True, "color": GOLD, "space_after": 4}),
                (action, {"size": 14, "bold": True, "color": BURGUNDY, "space_after": 0}),
            ],
        )

    set_notes(
        slide,
        """Keep ownership simple: Sponsor, Lead, Supporting Ministries.
Roadmap is deliberately light—four milestones, one year.
Ask who might serve as Operational Lead only if the room is ready.""",
    )


def slide_09_decision(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Recommendations & Decision", "What We Ask Leadership To Do")
    add_footer(slide, 9)

    recs = [
        ("Adopt", "Approve this Operational Workstream as the implementation approach for these Conference objectives."),
        ("Assign", "Identify an Executive Sponsor and Operational Lead within 90 days."),
        ("Plan", "Develop a shared annual communications plan with participating ministries."),
        ("Measure", "Track a short metric set and review progress quarterly."),
        ("Celebrate", "Publicly affirm ministry wins enabled by better coordination."),
    ]
    for i, (verb, text) in enumerate(recs):
        y = Inches(1.2) + i * Inches(1.05)
        card = add_round_rect(slide, Inches(0.5), y, Inches(12.3), Inches(0.95), WHITE, line=CARD_BORDER)
        verb_box = add_round_rect(slide, Inches(0.7), y + Inches(0.22), Inches(1.7), Inches(0.5), BURGUNDY if i < 2 else GOLD)
        shape_text(verb_box, verb, size=13, bold=True, color=WHITE if i < 2 else BURGUNDY_DARK)
        add_textbox(slide, Inches(2.7), y + Inches(0.25), Inches(9.8), Inches(0.5), text, size=14, color=CHARCOAL, anchor=MSO_ANCHOR.MIDDLE)

    set_notes(
        slide,
        """Seek a clear decision: endorse in principle; authorize Q1 lead identification.
If they want more detail later, offer a working session—not another full briefing.
Record the decision.""",
    )


def slide_10_closing(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, BURGUNDY_DARK)
    add_rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, GOLD)

    add_logo(slide, left=Inches(0.7), top=Inches(1.3), height=Inches(1.0))

    add_textbox(
        slide,
        Inches(0.7),
        Inches(2.8),
        Inches(12),
        Inches(0.7),
        '"Strong vision deserves faithful execution."',
        size=26,
        bold=True,
        color=GOLD,
        font_name="Georgia",
    )
    add_textbox(
        slide,
        Inches(0.7),
        Inches(3.8),
        Inches(11.8),
        Inches(1.6),
        "The Leadership Conference provided the vision.\n"
        "This workstream helps us execute it—with clearer collaboration,\n"
        "stronger communication, and support for every existing ministry.",
        size=16,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.7),
        Inches(6.3),
        Inches(11.8),
        Inches(0.4),
        "Ebenezer Baptist Church  ·  Making disciples and serving our community",
        size=12,
        color=GOLD_SOFT,
    )

    set_notes(
        slide,
        """Close with gratitude and the ask already stated on the prior slide.
Invite brief questions, then lock the decision.""",
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_01_title(prs)
    slide_02_the_point(prs)
    slide_03_build_upon(prs)
    slide_04_why(prs)
    slide_05_workstream(prs)
    slide_06_shared_service(prs)
    slide_07_outcomes_metrics(prs)
    slide_08_how_we_run(prs)
    slide_09_decision(prs)
    slide_10_closing(prs)

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    print(f"Slides: {TOTAL_SLIDES}")
    print(f"Logo: {OFFICIAL_LOGO if OFFICIAL_LOGO.exists() else 'missing'}")
    print(f"Brand: burgundy #802020 · gold #A08020")
