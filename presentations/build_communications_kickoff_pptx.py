#!/usr/bin/env python3
"""
EBC Communications Workstream — Kickoff / First 90 Days.

Small team, limited experience. AI consulted for goals, channels,
assessment — and for weekly drafts. People always decide.
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OFFICIAL_LOGO = ASSETS / "ebc-logo.png"
OUTPUT = ROOT / "EBC_Communications_Kickoff_90_Days.pptx"

PRES_TITLE = "Communications Workstream — Kickoff & First 90 Days"

BURGUNDY = RGBColor(0x80, 0x20, 0x20)
BURGUNDY_DARK = RGBColor(0x50, 0x14, 0x18)
GOLD = RGBColor(0xA0, 0x80, 0x20)
GOLD_SOFT = RGBColor(0xC4, 0xA8, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF2, 0xF0)
MID_GRAY = RGBColor(0xD8, 0xD2, 0xCC)
CHARCOAL = RGBColor(0x2A, 0x2A, 0x2A)
SOFT_CHARCOAL = RGBColor(0x5C, 0x55, 0x52)
CARD_BORDER = RGBColor(0xE2, 0xDC, 0xD6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_SLIDES = 11


def _set_run_font(run, *, size: Pt, bold: bool = False, color: RGBColor = CHARCOAL, name: str = "Calibri") -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _set_tf_margins(tf, margin: float = 0.08) -> None:
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.6)
    tf.margin_bottom = Inches(margin * 0.6)


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = CHARCOAL,
    align=PP_ALIGN.LEFT,
    font_name: str = "Calibri",
    anchor=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=Pt(size), bold=bold, color=color, name=font_name)


def add_rect(slide, left, top, width, height, fill: RGBColor, *, line: RGBColor | None = None, line_pt: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def add_round_rect(slide, left, top, width, height, fill: RGBColor, *, line: RGBColor | None = None, line_pt: float = 1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def shape_text(shape, text: str, *, size: int = 12, bold: bool = False, color: RGBColor = WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    _set_tf_margins(tf)
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=Pt(size), bold=bold, color=color)


def shape_multiline(shape, lines: Sequence[Tuple[str, dict]], *, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    _set_tf_margins(tf, 0.12)
    try:
        tf._txBody.bodyPr.set(
            qn("a:anchor"),
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    for i, (text, opts) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opts.get("space_after", 2))
        run = p.add_run()
        run.text = text
        _set_run_font(
            run,
            size=Pt(opts.get("size", 12)),
            bold=opts.get("bold", False),
            color=opts.get("color", WHITE),
        )


def add_logo(slide, *, left=Inches(0.4), top=Inches(0.22), height=Inches(0.55)) -> None:
    if OFFICIAL_LOGO.exists():
        slide.shapes.add_picture(str(OFFICIAL_LOGO), left, top, height=height)


def add_footer(slide, page: int) -> None:
    add_rect(slide, Inches(0.45), Inches(7.08), Inches(12.43), Pt(1.5), GOLD)
    add_textbox(slide, Inches(0.45), Inches(7.15), Inches(9.2), Inches(0.25), PRES_TITLE, size=9, color=SOFT_CHARCOAL)
    add_textbox(
        slide,
        Inches(9.8),
        Inches(7.15),
        Inches(3.1),
        Inches(0.25),
        f"Ebenezer Baptist Church  |  {page}",
        size=9,
        color=SOFT_CHARCOAL,
        align=PP_ALIGN.RIGHT,
    )


def add_slide_header(slide, title: str, eyebrow: str | None = None) -> None:
    add_logo(slide)
    title_left = Inches(3.35)
    if eyebrow:
        add_textbox(slide, title_left, Inches(0.2), Inches(9.4), Inches(0.25), eyebrow.upper(), size=10, bold=True, color=GOLD)
        add_textbox(slide, title_left, Inches(0.42), Inches(9.4), Inches(0.45), title, size=24, bold=True, color=BURGUNDY)
    else:
        add_textbox(slide, title_left, Inches(0.28), Inches(9.4), Inches(0.5), title, size=24, bold=True, color=BURGUNDY)
    add_rect(slide, title_left, Inches(0.95), Inches(1.3), Pt(3), BURGUNDY)


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_01_title(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, BURGUNDY_DARK)
    add_rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, GOLD)
    add_logo(slide, left=Inches(0.7), top=Inches(1.3), height=Inches(1.05))

    add_textbox(slide, Inches(0.7), Inches(2.7), Inches(11), Inches(0.35), "STEP 2  ·  AFTER ENDORSEMENT", size=12, bold=True, color=GOLD)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(3.15),
        Inches(12),
        Inches(1.1),
        "Communications Kickoff\n& First 90 Days",
        size=34,
        bold=True,
        color=WHITE,
    )
    add_rect(slide, Inches(0.7), Inches(4.5), Inches(1.5), Pt(2.5), GOLD)
    add_textbox(
        slide,
        Inches(0.7),
        Inches(4.75),
        Inches(11.5),
        Inches(1.0),
        "Small team. Limited experience.\nWe consult AI to set goals, choose channels, draft, and assess.",
        size=16,
        color=GOLD_SOFT,
    )
    add_textbox(slide, Inches(0.7), Inches(6.5), Inches(11), Inches(0.35), "Ebenezer Baptist Church", size=12, color=MID_GRAY)

    set_notes(
        slide,
        """Open: Endorsement is done. This kickoff stands up the operating rhythm.
Lead with scarcity + AI consulting for goals, channels, drafts, AND assessment — not drafts only.""",
    )


def slide_02_the_point(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "The Point", "What We Are Starting")
    add_footer(slide, 2)

    cards = [
        ("Tiny team", "Lead + 1–2 helpers.\nNot a department."),
        ("Consult AI", "Goals · Channels ·\nDrafts · Assess"),
        ("Weekly rhythm", "Same short meeting.\nChecklist every week."),
        ("People decide", "Lead OKs ordinary.\nPastor if sensitive."),
    ]
    for i, (t, d) in enumerate(cards):
        x = Inches(0.5) + i * Inches(3.2)
        add_round_rect(slide, x, Inches(1.4), Inches(3.05), Inches(3.5), WHITE, line=CARD_BORDER)
        add_rect(slide, x, Inches(1.4), Inches(3.05), Inches(0.12), BURGUNDY if i % 2 == 0 else GOLD)
        num = add_round_rect(slide, x + Inches(1.1), Inches(1.75), Inches(0.85), Inches(0.85), BURGUNDY)
        shape_text(num, str(i + 1), size=20, bold=True, color=WHITE)
        add_textbox(slide, x + Inches(0.2), Inches(2.85), Inches(2.65), Inches(0.5), t, size=16, bold=True, color=BURGUNDY, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.2), Inches(3.45), Inches(2.65), Inches(1.1), d, size=13, color=SOFT_CHARCOAL, align=PP_ALIGN.CENTER)

    banner = add_round_rect(slide, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.4), LIGHT_GRAY, line=GOLD, line_pt=1.25)
    shape_multiline(
        banner,
        [
            ("Why AI is heavy", {"size": 12, "bold": True, "color": GOLD, "space_after": 6}),
            (
                "We don’t have professional communicators — so we consult AI like an expert advisor, then humans approve.",
                {"size": 15, "bold": True, "color": BURGUNDY, "space_after": 0},
            ),
        ],
        align=PP_ALIGN.CENTER,
    )

    set_notes(
        slide,
        """Card 2 is the key correction: AI is not only a writer — we consult it for goals, channels, drafts, and assessment.
Banner explains the scarcity rationale in one sentence.""",
    )


def slide_03_consult_ai(prs: Presentation) -> None:
    """Core slide: Goals / Channels / Assess (+ weekly execute)."""
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Consult AI Like an Advisor", "Goals · Channels · Assess")
    add_footer(slide, 3)

    pillars = [
        (
            "GOALS",
            "What are we trying to achieve?",
            [
                "Clarity for guests & members",
                "Fewer conflicting announcements",
                "Pastor interrupted less",
                "Consistent public presence",
            ],
            "Ask AI to propose 90-day goals.\nLead + Sponsor choose.",
        ),
        (
            "CHANNELS",
            "Where should we show up?",
            [
                "Bulletin / in-service cues",
                "Website & email notes",
                "Facebook / YouTube messaging",
                "What to pause or skip",
            ],
            "Ask AI which channels fit\nour capacity — not every platform.",
        ),
        (
            "ASSESS",
            "Did it work? What’s next?",
            [
                "What shipped on time?",
                "Where details conflicted?",
                "Which channels were quiet?",
                "What to prioritize next week?",
            ],
            "Ask AI to score the week\nand recommend next focus.",
        ),
    ]

    for i, (title, question, items, footer) in enumerate(pillars):
        x = Inches(0.45) + i * Inches(4.25)
        card = add_round_rect(slide, x, Inches(1.25), Inches(4.1), Inches(5.3), WHITE, line=CARD_BORDER)
        add_rect(slide, x, Inches(1.25), Inches(4.1), Inches(1.15), BURGUNDY if i != 1 else BURGUNDY_DARK)
        add_textbox(slide, x + Inches(0.15), Inches(1.35), Inches(3.8), Inches(0.4), title, size=18, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.15), Inches(1.8), Inches(3.8), Inches(0.45), question, size=12, color=WHITE, align=PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            iy = Inches(2.6) + j * Inches(0.55)
            pill = add_round_rect(slide, x + Inches(0.2), iy, Inches(3.7), Inches(0.48), LIGHT_GRAY)
            shape_text(pill, item, size=12, bold=True, color=CHARCOAL)
        add_textbox(slide, x + Inches(0.2), Inches(5.0), Inches(3.7), Inches(1.2), footer, size=12, bold=True, color=BURGUNDY, align=PP_ALIGN.CENTER)

    set_notes(
        slide,
        """This is the heart of the kickoff. AI is consulted three ways before/around weekly drafting:
GOALS — set direction for 90 days (and refresh monthly).
CHANNELS — choose where a small team can actually show up.
ASSESS — weekly/monthly scorecard so we learn without a professional analyst.
Always: AI proposes; Lead/Sponsor dispose; Pastor for sensitive.""",
    )


def slide_04_full_loop(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "The Full Loop", "Discover → Execute → Evaluate")
    add_footer(slide, 4)

    stages = [
        ("Discover", "Consult AI", "What needs saying?\nWhich goals & channels\nthis week?"),
        ("Execute", "AI drafts", "One message package:\nbulletin · web · social\n· Media slide copy"),
        ("Evaluate", "Assess", "What worked?\nWhat broke?\nWhat next?"),
    ]
    for i, (name, mode, body) in enumerate(stages):
        x = Inches(0.7) + i * Inches(4.15)
        box = add_round_rect(slide, x, Inches(1.6), Inches(3.8), Inches(3.2), BURGUNDY if i != 1 else GOLD)
        text_c = WHITE if i != 1 else BURGUNDY_DARK
        shape_multiline(
            box,
            [
                (name, {"size": 20, "bold": True, "color": text_c, "space_after": 6}),
                (mode, {"size": 13, "bold": True, "color": GOLD if i != 1 else BURGUNDY, "space_after": 10}),
                (body, {"size": 14, "color": text_c, "space_after": 0}),
            ],
        )
        if i < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(3.85), Inches(2.9), Inches(0.28), Inches(0.45))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GOLD if i == 0 else BURGUNDY
            arrow.line.fill.background()

    # feedback arrow note
    banner = add_round_rect(slide, Inches(0.7), Inches(5.2), Inches(11.9), Inches(1.35), LIGHT_GRAY, line=CARD_BORDER)
    shape_multiline(
        banner,
        [
            ("Evaluate feeds the next Discover", {"size": 16, "bold": True, "color": BURGUNDY, "space_after": 6}),
            (
                "Assessment is not a report for the shelf — it tells AI and the team what to prioritize next week.",
                {"size": 14, "color": CHARCOAL, "space_after": 0},
            ),
        ],
        align=PP_ALIGN.CENTER,
    )

    set_notes(
        slide,
        """Map the three consult modes into the operating loop.
Discover = goals + channels for the week.
Execute = drafts (what people already expect from AI).
Evaluate = assess — closes the loop.
Without Assess, AI only writes; with Assess, AI helps the team get smarter.""",
    )


def slide_05_tiny_team(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "The Tiny Team", "Roles That Fit Real Life")
    add_footer(slide, 5)

    roles = [
        ("Executive Sponsor", "Pastor (or designee)", "Air cover. Sensitive approvals. Affirms goals."),
        ("Operational Lead", "Reliable organizer", "Runs the loop. Consults AI. Final ordinary OK."),
        ("Helper(s)", "1–2 volunteers", "Help gather inputs; post approved packages."),
        ("AI Advisor", "Expert teammate", "Goals · channels · drafts · weekly assessment."),
    ]
    for i, (role, who, what) in enumerate(roles):
        y = Inches(1.25) + i * Inches(1.3)
        add_round_rect(slide, Inches(0.5), y, Inches(12.3), Inches(1.15), WHITE, line=CARD_BORDER)
        badge = add_round_rect(
            slide,
            Inches(0.7),
            y + Inches(0.25),
            Inches(2.6),
            Inches(0.65),
            BURGUNDY if i < 2 else (GOLD if i == 3 else LIGHT_GRAY),
        )
        shape_text(badge, role, size=12, bold=True, color=WHITE if i != 2 else BURGUNDY)
        add_textbox(slide, Inches(3.6), y + Inches(0.2), Inches(3.2), Inches(0.35), who, size=14, bold=True, color=BURGUNDY)
        add_textbox(slide, Inches(3.6), y + Inches(0.55), Inches(8.9), Inches(0.4), what, size=13, color=SOFT_CHARCOAL)

    set_notes(
        slide,
        """AI Advisor row must say goals/channels/drafts/assessment — not “writer” only.
Lead’s job is to run the consult loop, not invent strategy from scratch.""",
    )


def slide_06_weekly_rhythm(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Weekly Meeting Agenda", "45 Minutes, Same Every Week")
    add_footer(slide, 6)

    agenda = [
        ("10 min", "Assess last week", "Ask AI: what shipped, what conflicted, what’s quiet?"),
        ("10 min", "Goals & channels", "Ask AI: what matters this week, where do we post?"),
        ("15 min", "Shape & clear", "AI drafts package → team edits → Lead/Pastor OK"),
        ("10 min", "Send & assign", "Who posts what; Media gets slide copy"),
    ]
    for i, (time, title, desc) in enumerate(agenda):
        y = Inches(1.25) + i * Inches(1.3)
        card = add_round_rect(slide, Inches(0.5), y, Inches(12.3), Inches(1.15), WHITE, line=CARD_BORDER)
        time_box = add_round_rect(slide, Inches(0.7), y + Inches(0.25), Inches(1.5), Inches(0.65), GOLD if i == 0 else BURGUNDY)
        shape_text(time_box, time, size=13, bold=True, color=BURGUNDY_DARK if i == 0 else WHITE)
        add_textbox(slide, Inches(2.5), y + Inches(0.2), Inches(9.9), Inches(0.35), title, size=16, bold=True, color=BURGUNDY)
        add_textbox(slide, Inches(2.5), y + Inches(0.55), Inches(9.9), Inches(0.4), desc, size=13, color=SOFT_CHARCOAL)

    set_notes(
        slide,
        """Agenda deliberately starts with Assess — so consulting AI for evaluation is habitual.
Then goals/channels for this week, then drafts, then send.
Gold on Assess to draw the eye.""",
    )


def slide_07_approvals(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Who Approves What", "AI Never Has the Final Word")
    add_footer(slide, 7)

    rows = [
        ("90-day goals", "AI proposes options", "Lead + Sponsor choose"),
        ("Channel focus", "AI recommends mix", "Lead confirms (capacity)"),
        ("Weekly package", "AI drafts copy", "Lead OKs ordinary"),
        ("Sensitive items", "AI may draft", "Pastor must approve"),
        ("Weekly assess", "AI scores & suggests", "Lead decides next focus"),
    ]
    add_rect(slide, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.5), BURGUNDY)
    add_textbox(slide, Inches(0.7), Inches(1.35), Inches(3.5), Inches(0.35), "Decision", size=13, bold=True, color=WHITE)
    add_textbox(slide, Inches(4.4), Inches(1.35), Inches(4.0), Inches(0.35), "AI role", size=13, bold=True, color=WHITE)
    add_textbox(slide, Inches(8.6), Inches(1.35), Inches(4.0), Inches(0.35), "Human role", size=13, bold=True, color=WHITE)

    for i, (decision, ai, human) in enumerate(rows):
        y = Inches(1.85) + i * Inches(0.9)
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        add_rect(slide, Inches(0.5), y, Inches(12.3), Inches(0.85), bg, line=CARD_BORDER)
        add_textbox(slide, Inches(0.7), y + Inches(0.25), Inches(3.5), Inches(0.4), decision, size=14, bold=True, color=BURGUNDY)
        add_textbox(slide, Inches(4.4), y + Inches(0.25), Inches(4.0), Inches(0.4), ai, size=13, color=CHARCOAL)
        add_textbox(slide, Inches(8.6), y + Inches(0.25), Inches(4.0), Inches(0.4), human, size=13, bold=True, color=CHARCOAL)

    set_notes(
        slide,
        """Show AI involved in goals, channels, drafts, AND assess — with humans always deciding.
Pastor row is the trust builder.""",
    )


def slide_08_partners(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "How Partners Work", "Submit — Don’t Freelance")
    add_footer(slide, 8)

    partners = [
        ("Ministries", "Send requests by the deadline.\nOne version. No surprise Sunday asks."),
        ("Church Office", "Partner on publishing.\nNot the only writer for everything."),
        ("Media Ministry", "Owns AV: sound, stream, slides.\nCommunications brings the words."),
        ("AI Advisor", "Consulted on goals, channels,\ndrafts, and weekly assessment."),
    ]
    for i, (t, d) in enumerate(partners):
        r, c = i // 2, i % 2
        x = Inches(0.5) + c * Inches(6.4)
        y = Inches(1.35) + r * Inches(2.5)
        add_round_rect(slide, x, y, Inches(6.15), Inches(2.25), WHITE, line=CARD_BORDER)
        add_rect(slide, x, y, Inches(6.15), Inches(0.55), GOLD if i == 3 else (BURGUNDY if i % 2 == 0 else BURGUNDY_DARK))
        add_textbox(slide, x + Inches(0.25), y + Inches(0.12), Inches(5.65), Inches(0.4), t, size=16, bold=True, color=BURGUNDY_DARK if i == 3 else WHITE)
        add_textbox(slide, x + Inches(0.3), y + Inches(0.85), Inches(5.55), Inches(1.1), d, size=15, color=CHARCOAL)

    set_notes(
        slide,
        """AI listed as a partner so leadership sees consulting is structural.
Media vs Communications boundary unchanged.""",
    )


def slide_09_ninety_days(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "First 90 Days", "Goals → Rhythm → Assess")
    add_footer(slide, 9)

    add_rect(slide, Inches(0.7), Inches(3.35), Inches(11.9), Pt(4), GOLD)

    phases = [
        (
            "Days 1–30",
            "Set direction",
            [
                "Name Lead + helpers",
                "Consult AI on 90-day goals",
                "Consult AI on channel mix",
                "Lock weekly meeting",
            ],
        ),
        (
            "Days 31–60",
            "Run the loop",
            [
                "Discover → Execute weekly",
                "Assess every meeting",
                "Same details everywhere",
                "Pastor only when needed",
            ],
        ),
        (
            "Days 61–90",
            "Prove & adjust",
            [
                "AI-assisted scorecard",
                "Keep / cut channels",
                "Refresh goals if needed",
                "Report to Sponsor",
            ],
        ),
    ]
    for i, (title, subtitle, items) in enumerate(phases):
        x = Inches(0.7) + i * Inches(4.1)
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(1.45), Inches(3.15), Inches(0.45), Inches(0.45))
        marker.fill.solid()
        marker.fill.fore_color.rgb = BURGUNDY
        marker.line.fill.background()
        shape_text(marker, str(i + 1), size=12, bold=True, color=WHITE)

        card = add_round_rect(slide, x, Inches(1.25), Inches(3.85), Inches(1.65), WHITE, line=CARD_BORDER)
        shape_multiline(
            card,
            [
                (title, {"size": 13, "bold": True, "color": GOLD, "space_after": 4}),
                (subtitle, {"size": 17, "bold": True, "color": BURGUNDY, "space_after": 0}),
            ],
        )
        for j, item in enumerate(items):
            iy = Inches(3.9) + j * Inches(0.55)
            pill = add_round_rect(slide, x, iy, Inches(3.85), Inches(0.48), LIGHT_GRAY)
            shape_text(pill, item, size=12, bold=True, color=CHARCOAL)

    set_notes(
        slide,
        """Month 1 explicitly consults AI for goals and channels before the rhythm hardens.
Month 2 runs Discover→Execute→Evaluate.
Month 3 uses assessment to keep/cut channels and refresh goals.""",
    )


def slide_10_success(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Success at Day 90", "Direction + Rhythm + Learning")
    add_footer(slide, 10)

    checks = [
        ("Goals chosen with AI input", "Written 90-day goals the Sponsor affirmed"),
        ("Channels match capacity", "We show up where we can — not everywhere"),
        ("Weekly assess habit", "Every meeting starts with last week’s score"),
        ("Fewer conflicting announcements", "Same details across pulpit, bulletin, web, social"),
        ("Pastor interrupted less", "Requests hit the deadline"),
        ("Team not drowning", "AI consulted; checklist holds; humans still decide"),
    ]
    for i, (t, d) in enumerate(checks):
        r, c = i // 2, i % 2
        x = Inches(0.5) + c * Inches(6.4)
        y = Inches(1.25) + r * Inches(1.7)
        add_round_rect(slide, x, y, Inches(6.15), Inches(1.5), WHITE, line=CARD_BORDER)
        add_rect(slide, x, y, Pt(5), Inches(1.5), GOLD)
        add_textbox(slide, x + Inches(0.3), y + Inches(0.25), Inches(5.6), Inches(0.4), t, size=14, bold=True, color=BURGUNDY)
        add_textbox(slide, x + Inches(0.3), y + Inches(0.75), Inches(5.6), Inches(0.5), d, size=13, color=SOFT_CHARCOAL)

    set_notes(
        slide,
        """Success includes consulting AI for goals/channels/assess — not only “we posted more.”
Learning loop is the proof the scarcity model works.""",
    )


def slide_11_decide(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WHITE)
    add_slide_header(slide, "Decisions Needed Now", "Lock the First 90 Days")
    add_footer(slide, 11)

    asks = [
        ("1", "Name Lead + helpers", "Reliable organizers — résumés optional."),
        ("2", "Schedule AI goal session", "Consult AI on 90-day goals; Sponsor affirms."),
        ("3", "Schedule AI channel session", "Choose where we will (and won’t) show up."),
        ("4", "Set weekly loop meeting", "Assess → Goals/Channels → Draft → Send."),
        ("5", "Confirm approval path", "Lead ordinary · Pastor sensitive · AI never final."),
    ]
    for i, (num, title, desc) in enumerate(asks):
        y = Inches(1.2) + i * Inches(1.05)
        add_round_rect(slide, Inches(0.5), y, Inches(12.3), Inches(0.95), WHITE, line=CARD_BORDER)
        badge = add_round_rect(slide, Inches(0.7), y + Inches(0.2), Inches(0.7), Inches(0.55), BURGUNDY)
        shape_text(badge, num, size=16, bold=True, color=WHITE)
        add_textbox(slide, Inches(1.7), y + Inches(0.15), Inches(10.8), Inches(0.35), title, size=16, bold=True, color=BURGUNDY)
        add_textbox(slide, Inches(1.7), y + Inches(0.5), Inches(10.8), Inches(0.35), desc, size=13, color=SOFT_CHARCOAL)

    set_notes(
        slide,
        """Decisions 2 and 3 are the new explicit asks — AI goal session and AI channel session.
Weekly meeting agenda always includes Assess.
Leave with names and calendar holds.""",
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_01_title(prs)
    slide_02_the_point(prs)
    slide_03_consult_ai(prs)
    slide_04_full_loop(prs)
    slide_05_tiny_team(prs)
    slide_06_weekly_rhythm(prs)
    slide_07_approvals(prs)
    slide_08_partners(prs)
    slide_09_ninety_days(prs)
    slide_10_success(prs)
    slide_11_decide(prs)

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out}")
    print(f"Slides: {TOTAL_SLIDES}")
